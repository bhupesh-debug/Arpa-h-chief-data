"""10-slide deck, built directly with python-pptx (no pptx skill installed)."""
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title, bullets, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    tx = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9)).text_frame
    tx.word_wrap = True
    p = tx.paragraphs[0]; p.text = title; p.font.size = Pt(30); p.font.bold = True
    if subtitle:
        tx2 = s.shapes.add_textbox(Inches(0.6), Inches(1.15), Inches(12.1), Inches(0.6)).text_frame
        tx2.word_wrap = True
        p = tx2.paragraphs[0]; p.text = subtitle; p.font.size = Pt(16); p.font.italic = True
        top = 1.9
    else:
        top = 1.5
    tx3 = s.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(5.2)).text_frame
    tx3.word_wrap = True
    for i, b in enumerate(bullets):
        p = tx3.paragraphs[0] if i == 0 else tx3.add_paragraph()
        p.text = "•  " + b; p.font.size = Pt(19); p.space_after = Pt(10)
    return s


slide("Executive Data Leadership for a Health Research Agency",
      ["A working portfolio of the artifacts a Chief of IT and Data Technology governs",
       "Responsible-AI routing  •  Lakehouse quality gates  •  ATO compliance tracking",
       "PORTFOLIO DEMONSTRATION on synthetic data — NOT evidence of a C-suite/SES role,",
       "federal employment, a clearance, or any executive record."],
      subtitle="JD-to-ML-Project build  •  ARPA-H Chief of IT and Data Technology (JD-based)")

slide("Problem Overview",
      ["A new agency must stand up enterprise technology, data, AI, and cybersecurity at once —",
       "while every decision is auditable, every dataset is governed, and every system is authorized.",
       "One-liner: this project shows HOW a chief would run that function, in working code,",
       "not in slides alone."])

slide("JD Breakdown",
      ["Must-have: 10+ years senior executive technology leadership; 5 years C-suite/SES/equivalent",
       "AI strategy + responsible-AI governance; enterprise data architecture; FISMA / FedRAMP / NIST",
       "Software engineering leadership (DevSecOps, agile, CI/CD); cybersecurity and A&A ownership",
       "Hidden expectations: build the function from near-zero; advise the Director; move at startup speed",
       "inside federal constraints; own acquisition and vendor performance."])

slide("ML Approach — Honest Fit",
      ["This is a chief-level leadership role, NOT a hands-on ML job.",
       "A generic model demo would read as tone-deaf and would misrepresent the candidate.",
       "So the artifact is executive infrastructure: transparent RULES, not black boxes.",
       "Risk scores, quality gates, and compliance scores an executive can defend in a hearing.",
       "No deep learning, no LLM demo, none claimed."])

slide("Data Strategy",
      ["All data synthetic, seeded, and labeled fictional — no real agency data anywhere.",
       "14 AI use cases with sensitivity / autonomy / impact / opacity / deployment attributes",
       "4 research-style datasets: 243 trials, 600 biosamples, 400 images, 120 grants",
       "8 systems x 30 NIST-style controls = 240 assessments, 115 POA&M items",
       "Injected defects are disclosed: dirty trials, broken image references, one veto scenario."])

slide("Solution Design",
      ["1. AI governance engine: weighted risk score (0-100), 4 tiers, approval routes with SLAs,",
       "   policy floors (restricted data + autonomy can never score below High).",
       "2. Lakehouse gates: completeness, uniqueness, range, reference, allowed-value checks;",
       "   Bronze / Silver / Gold promotion — a failed reference check caps at Bronze.",
       "3. ATO tracker: control scores per system; open critical POA&M VETOES authorization."])

slide("Results & Insights",
      ["AI routing: Low 3, Moderate 5, High 4, Critical 2 — floor rule fires once (AI-10).",
       "Lakehouse: grants Gold (1.000), biosamples Silver (0.833), imaging Bronze (0.500),",
       "trials Bronze (0.400) — every tier demonstrated, every defect caught by a named check.",
       "ATO: 2 Authorized, 6 Not authorized. Identity Platform scores 0.95 yet is NOT authorized:",
       "one open critical POA&M vetoes it. That veto IS the governance point."])

slide("Business Impact",
      ["Governance speed: tiered SLAs (5 / 15 / 30 / 45 days) replace ad-hoc approvals.",
       "Trust: every routing, promotion, and authorization decision is reproducible from rules.",
       "Risk reduction: critical gaps cannot hide inside a high average score.",
       "For a chief: these are the operating rhythms of the function, not a model accuracy claim."])

slide("Future Improvements",
      ["Port the batch pipeline to a Databricks-style job (Delta tables, scheduled workflows).",
       "Add a human-appeals path and audit trail for contested AI routing decisions.",
       "Add dataset lineage and drift monitoring to the lakehouse gates.",
       "Add POA&M aging dashboards and overdue-escalation rules to the ATO tracker.",
       "Re-validate every threshold on real agency data before any operational use."])

slide("About This Project",
      ["Created by Bhupesh Chandra Dimri  •  bhupesh7cd@gmail.com",
       "GitHub link: [placeholder — no repository created for this build]",
       "Portfolio demonstration only. Not evidence of executive, federal, or clearance status.",
       "All data synthetic. All limitations are stated in the project README."])

prs.save("/home/hatch/workspace/ml-project-arpa-h-chief-data/reports/deck.pptx")
print("deck slides:", len(prs.slides.__iter__.__self__._sldIdLst))
